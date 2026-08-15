#!/usr/bin/env python3
"""Create a minimal, loss-checked PNN project from DOCX, TXT, or Markdown.

This script performs no semantic analysis. It only extracts text, finds source
chapter headings, writes ordered PNN slices, and initializes project state.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import sys
from datetime import datetime, timezone
from pathlib import Path


CHAPTER_RE = re.compile(
    r"(?im)^[ \t]*Chapter[ \t]+(?P<label>\d+|[IVXLCDM]+|[A-Za-z]+)\b[^\r\n]*"
)
SUPPORTED_SUFFIXES = {".docx", ".txt", ".md", ".markdown"}


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_text(text: str) -> str:
    return sha256_bytes(text.encode("utf-8"))


def read_plain_text(path: Path) -> tuple[str, str]:
    raw = path.read_bytes()
    if raw.startswith(b"\xef\xbb\xbf"):
        return raw.decode("utf-8-sig"), "plain-text:utf-8-sig"
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16"), "plain-text:utf-16"
    try:
        return raw.decode("utf-8"), "plain-text:utf-8"
    except UnicodeDecodeError as exc:
        raise ValueError(
            "文本不是 UTF-8/UTF-16。请先另存为 UTF-8，避免静默损坏原文。"
        ) from exc


def read_docx_text(path: Path) -> tuple[str, str]:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError(
            "读取 DOCX 需要 python-docx：python -m pip install python-docx"
        ) from exc

    doc = Document(path)
    nonempty_tables = []
    for table_index, table in enumerate(doc.tables, start=1):
        cell_text = "".join(cell.text for row in table.rows for cell in row.cells)
        if cell_text.strip():
            nonempty_tables.append(table_index)
    if nonempty_tables:
        raise ValueError(
            "检测到含正文的 DOCX 表格。为避免改变原文顺序，本最小工具拒绝静默提取；"
            "请先另存为纯文本后再运行。表格编号："
            + ", ".join(map(str, nonempty_tables))
        )

    # DOCX does not store a canonical plain-text byte stream. Paragraph text is
    # joined with LF, while the untouched DOCX copy remains the source authority.
    return "\n".join(paragraph.text for paragraph in doc.paragraphs), "docx:paragraphs-lf"


def read_source(path: Path) -> tuple[str, str]:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(f"不支持的输入格式：{suffix or '<无扩展名>'}")
    if suffix == ".docx":
        return read_docx_text(path)
    return read_plain_text(path)


def choose_pnn(label: str, fallback: int, used: set[int]) -> int:
    if label.isdigit():
        number = int(label)
        if 0 <= number <= 999 and number not in used:
            return number
    while fallback in used or fallback == 0:
        fallback += 1
    return fallback


def split_chapters(text: str) -> list[dict[str, object]]:
    matches = list(CHAPTER_RE.finditer(text))
    if not matches:
        return [
            {
                "id": "P01",
                "source_heading": "未检测到 Chapter 标题",
                "start": 0,
                "end": len(text),
                "text": text,
            }
        ]

    chapters: list[dict[str, object]] = []
    used: set[int] = set()
    fallback = 1

    first_start = matches[0].start()
    preamble = text[:first_start]
    first_chapter_starts_at = first_start
    if preamble.strip():
        chapters.append(
            {
                "id": "P00",
                "source_heading": "正文前内容",
                "start": 0,
                "end": first_start,
                "text": preamble,
            }
        )
        used.add(0)
    elif first_start > 0:
        # Keep formatting-only leading whitespace without inventing a P00.
        first_chapter_starts_at = 0

    for index, match in enumerate(matches):
        start = first_chapter_starts_at if index == 0 else match.start()
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        number = choose_pnn(match.group("label"), fallback, used)
        used.add(number)
        fallback = max(fallback, number + 1)
        chapters.append(
            {
                "id": f"P{number:02d}",
                "source_heading": match.group(0).strip(),
                "start": start,
                "end": end,
                "text": text[start:end],
            }
        )

    return chapters


def write_text_exact(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        handle.write(text)


def read_text_exact(path: Path) -> str:
    with path.open("r", encoding="utf-8", newline="") as handle:
        return handle.read()


def write_json(path: Path, value: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="\n") as handle:
        json.dump(value, handle, ensure_ascii=False, indent=2)
        handle.write("\n")


def render_progress(title: str, chapter_ids: list[str]) -> str:
    lines = [
        "# 制作进度",
        "",
        f"**项目**：{title}",
        f"**章节数**：{len(chapter_ids)}",
        "**已通过**：0",
        f"**下一章**：{chapter_ids[0] if chapter_ids else '无'}",
        "",
        "## 章节状态",
        "",
    ]
    lines.extend(f"- [ ] {chapter_id}：待处理" for chapter_id in chapter_ids)
    lines.append("")
    return "\n".join(lines)


def initialize_project(source: Path, output: Path, title: str) -> dict[str, object]:
    if output.exists() and any(output.iterdir()):
        raise FileExistsError(f"目标目录不是空目录，拒绝覆盖：{output}")
    output.mkdir(parents=True, exist_ok=True)

    canonical_text, extraction_method = read_source(source)
    chapters = split_chapters(canonical_text)
    if not chapters:
        raise ValueError("没有得到任何章节。")

    source_dir = output / "source"
    chapter_dir = source_dir / "chapters"
    work_dir = output / "work"
    deliverables_dir = output / "deliverables" / "chapters"
    chapter_dir.mkdir(parents=True, exist_ok=True)
    work_dir.mkdir(parents=True, exist_ok=True)
    deliverables_dir.mkdir(parents=True, exist_ok=True)

    original_copy = source_dir / f"original.input{source.suffix.lower()}"
    shutil.copy2(source, original_copy)

    manifest_chapters = []
    for order, chapter in enumerate(chapters, start=1):
        chapter_id = str(chapter["id"])
        chapter_text = str(chapter["text"])
        chapter_path = chapter_dir / f"{chapter_id}.txt"
        write_text_exact(chapter_path, chapter_text)
        manifest_chapters.append(
            {
                "id": chapter_id,
                "order": order,
                "source_heading": chapter["source_heading"],
                "start_char": chapter["start"],
                "end_char": chapter["end"],
                "characters": len(chapter_text),
                "sha256": sha256_text(chapter_text),
                "path": f"source/chapters/{chapter_id}.txt",
            }
        )

    reconstructed = "".join(
        read_text_exact(chapter_dir / f"{chapter['id']}.txt")
        for chapter in manifest_chapters
    )
    if reconstructed != canonical_text:
        raise RuntimeError("内部无损校验失败：章节重新拼接后与提取文本不一致。")

    source_bytes = source.read_bytes()
    now = datetime.now(timezone.utc).isoformat()
    manifest = {
        "version": 1,
        "title": title,
        "created_at": now,
        "source_filename": source.name,
        "source_sha256": sha256_bytes(source_bytes),
        "canonical_text_sha256": sha256_text(canonical_text),
        "canonical_characters": len(canonical_text),
        "extraction_method": extraction_method,
        "reconstruction_verified": True,
        "chapters": manifest_chapters,
    }
    write_json(source_dir / "manifest.json", manifest)

    state = {
        "version": 1,
        "title": title,
        "updated_at": now,
        "chapters": {
            chapter["id"]: {
                "status": "pending",
                "source_sha256": chapter["sha256"],
                "output_sha256": None,
            }
            for chapter in manifest_chapters
        },
    }
    write_json(work_dir / "state.json", state)
    chapter_ids = [str(chapter["id"]) for chapter in manifest_chapters]
    write_text_exact(work_dir / "制作进度.md", render_progress(title, chapter_ids))
    write_text_exact(
        work_dir / "项目记忆.md",
        "# 项目记忆\n\n> 由校验工具根据已通过章节的章末索引生成。首次处理暂无记录。\n",
    )
    write_text_exact(
        work_dir / "全文复盘输入.md",
        "# 全文复盘输入\n\n> 尚无通过校验的章节。\n",
    )
    return manifest


def main() -> int:
    parser = argparse.ArgumentParser(description="无损提取并拆分为 PNN 章节。")
    parser.add_argument("source", type=Path, help="DOCX、TXT 或 Markdown 原稿")
    parser.add_argument("--out", required=True, type=Path, help="新项目目录")
    parser.add_argument("--title", default=None, help="项目名；默认使用源文件名")
    args = parser.parse_args()

    source = args.source.resolve()
    if not source.is_file():
        parser.error(f"找不到原稿：{source}")
    title = args.title or source.stem

    try:
        manifest = initialize_project(source, args.out.resolve(), title)
    except Exception as exc:  # concise CLI boundary
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1

    ids = [chapter["id"] for chapter in manifest["chapters"]]
    print(f"PROJECT={args.out.resolve()}")
    print(f"CHAPTERS={','.join(ids)}")
    print(f"SOURCE_SHA256={manifest['source_sha256']}")
    print("RECONSTRUCTION=PASSED")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
